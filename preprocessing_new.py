"""
enriched_preprocessing_v3.py — Final clean version
Handles all NUST Bank Excel table patterns correctly.
"""
import re, json, hashlib
from pathlib import Path
from openpyxl import load_workbook

EXCEL_PATH       = "product_info.xlsx"
OUTPUT_JSONL     = "preprocessed_data/bank_data_master.jsonl"
RATE_SHEET_JSONL = "preprocessed_data/rate_sheet_only.jsonl"

ACCOUNT_NAME_MAP = {
    "LCA":"Little Champs Account","NAA":"NUST Asaan Account (NAA)","NWA":"NUST Waqaar Account",
    "PWRA":"PakWatan Remittance Account","RDA":"Roshan Digital Account","VPCA":"Value Plus Current Account",
    "VP-BA":"Value Plus Business Account","VPBA":"NUST Value Premium Business Account",
    "NSDA":"NUST Special Deposit Account","PLS":"Profit and Loss Sharing Account (PLS)",
    "CDA":"Current Deposit Account","NMA":"NUST Maximiser Account","NADA":"NUST Asaan Digital Account",
    "NADRA":"NUST Asaan Digital Remittance Account","NUST4Car":"NUST4Car Auto Finance",
    "ESFCA":"Exporters' Special Foreign Currency Account","NFDA":"NUST Freelancer Digital Account",
    "NSA":"NUST Sahar Accounts","PF":"NUST Personal Finance","NMC":"NUST Bank Mastercard",
    "NMF":"NUST Mortgage Finance","NSF":"NUST Sahar Finance","NIF":"NUST Imarat Finance",
    "NUF":"NUST Ujala Finance","NFMF":"NUST Flour Mill Finance","NFBF":"NUST Fauri Business Finance",
    "PMYB &ALS":"Prime Minister Youth Business & Agriculture Loan Scheme",
    "NRF":"NUST Rice Finance","NHF":"NUST Hunarmand Finance","Nust Life":"NUST Life Bancassurance Policy",
    "EFU Life":"EFU Life Bancassurance Policy","Jubilee Life ":"Jubilee Life Bancassurance Policy",
    "HOME REMITTANCE":"Home Remittance Services",
}
DEFAULT_NAME = "NUST Bank Product"
SKIP_SHEETS  = {"Main","Rate Sheet July 1 2024","Sheet1"}
PII_PATTERNS = {
    "CNIC":r'\b\d{5}-\d{7}-\d{1}\b',"PHONE":r'(\+92|0)3\d{2}-\d{7}|\b0\d{2}-\d{7,8}\b',
    "EMAIL":r'\b[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}\b',
    "IBAN":r'\bPK\d{2}[A-Z]{4}[A-Z0-9]{16}\b',"ACCOUNT_NUMBER":r'\b\d{10,14}\b',
}

def anonymize(t):
    for l,p in PII_PATTERNS.items(): t=re.sub(p,f"<{l}_MASKED>",t)
    return t

def clean_val(v):
    if v is None: return ""
    s=re.sub(r'\s+',' ',str(v).strip())
    try:
        f=float(s)
        if 0<f<1: return f"{f*100:.2f}%"
    except: pass
    return s

def sha(t): return hashlib.sha256(t.encode()).hexdigest()
def is_noise(v): return not v or v in ('–','-','nan','None','*')
def is_pct(v): return '%' in str(v)

def is_question(t):
    t=t.strip()
    if t.endswith('?'): return True
    return t.lower().startswith(('what','how','why','when','where','can','is','are',
                                  'does','do','which','who','will','would','should'))

def read_sparse(ws):
    skip=set()
    for m in ws.merged_cells.ranges:
        for r in range(m.min_row,m.max_row+1):
            for c in range(m.min_col,m.max_col+1):
                if r!=m.min_row or c!=m.min_col: skip.add((r,c))
    data={}
    for r in range(1,ws.max_row+1):
        for c in range(1,ws.max_column+1):
            if (r,c) in skip: continue
            v=clean_val(ws.cell(r,c).value)
            if v: data.setdefault(r,{})[c]=v
    return data

def rows_list(data):
    return [(r,data[r]) for r in sorted(data)]

def col_vals(cd): return [(c,v) for c,v in sorted(cd.items())]
def just_vals(cd): return [v for _,v in col_vals(cd)]

def is_table_row(cd):
    vals=[v for v in just_vals(cd) if not is_noise(v)]
    if len(vals)<2: return False
    if any(is_question(v) for v in vals): return False
    if any(is_pct(v) for v in vals): return True
    short=sum(1 for v in vals if len(v)<=80)
    return short/len(vals)>=0.6

# ── Table converters ─────────────────────────────────────────────────────────

def convert_multi_section(block_rows, question, account_name):
    if not block_rows: return ""
    first_row = col_vals(block_rows[0])
    label_col = first_row[0][0]
    sections  = [(c,v) for c,v in first_row if c!=label_col and not is_noise(v) and not is_question(v)]
    if not sections:
        return "\n".join(f"- {v}" for v in just_vals(block_rows[0]) if not is_noise(v))

    rightmost_col  = sections[-1][0]
    rightmost_name = sections[-1][1]
    main_sections  = [(c,n) for c,n in sections if c!=rightmost_col]

    def get_main_section(c):
        assigned=None
        for sc,sn in main_sections:
            if c>=sc: assigned=sn
        return assigned

    main_data   = {}
    sub_rows    = []
    current_sub = ""

    for cd in block_rows[1:]:
        cvs   = col_vals(cd)
        label = cd.get(label_col,'')
        if is_noise(label): label=''

        for c,v in cvs:
            if c==label_col or is_noise(v) or c>=rightmost_col: continue
            sec=get_main_section(c)
            if sec:
                if label not in main_data: main_data[label]={}
                main_data[label][sec]=v

        right=[(c,v) for c,v in cvs if c>=rightmost_col and not is_noise(v)]
        if right: sub_rows.append(right)

    lines=[]
    sec_names=[sn for _,sn in main_sections]
    if sec_names: lines.append(f"Comparison ({' vs '.join(sec_names)}):")
    for feat,sec_vals in main_data.items():
        if not feat: continue
        parts=[f"{sn}: {sec_vals[sn]}" for sn in sec_names if sn in sec_vals]
        lines.append(f"  • {feat}: {', '.join(parts)}" if parts else f"  • {feat}")

    if sub_rows:
        lines.append(f"\n{rightmost_name}:")
        for row_v in sub_rows:
            vals=[v for _,v in row_v]
            if len(vals)==1 and (vals[0].endswith(':') or not any(c.isdigit() for c in vals[0])):
                current_sub=vals[0].rstrip(':'); continue
            if any('tenure' in v.lower() or 'payout' in v.lower() for v in vals): continue
            if any(is_pct(v) for v in vals):
                rate_str=" | ".join(v for v in vals if not is_noise(v))
                prefix=f"{current_sub}: " if current_sub else ""
                lines.append(f"  • {prefix}{rate_str}")
            else:
                s=" | ".join(v for v in vals if not is_noise(v))
                if s:
                    if s.endswith(':'): current_sub=s.rstrip(':')
                    else: lines.append(f"  • {s}")

    return "\n".join(lines)

def convert_profit_inline(block_rows, question, account_name):
    if not block_rows: return ""
    first=col_vals(block_rows[0])
    cols={v.lower().strip():c for c,v in first}
    payment_col=next((c for k,c in cols.items() if 'profit payment' in k or 'payment' in k),None)
    rate_col   =next((c for k,c in cols.items() if 'profit rate' in k or ('rate' in k and 'payment' not in k)),None)
    tenure_col =next((c for k,c in cols.items() if 'tenure' in k or 'tenor' in k),None)
    payout_col =next((c for k,c in cols.items() if 'payout' in k),None)

    thresh=min(x for x in [payment_col,rate_col,tenure_col] if x is not None) if any([payment_col,rate_col,tenure_col]) else 999
    feature_lines=[]; profit_lines=[]

    for cd in block_rows:
        cvs=col_vals(cd)
        left =[v for c,v in cvs if not is_noise(v) and c<thresh]
        right={c:v for c,v in cvs if not is_noise(v) and c>=thresh}
        for v in left:
            v2=re.sub(r'^[o•\-\*]\s*','',v).strip()
            if v2 and v2 not in ('Profit Payment','Profit Rate'): feature_lines.append(v2)
        if payment_col and rate_col:
            pay=right.get(payment_col,''); rt=right.get(rate_col,'')
            if pay and rt and not is_noise(pay) and not is_noise(rt):
                profit_lines.append(f"Profit Payment: {pay}, Profit Rate: {rt}")
        elif tenure_col and rate_col:
            ten=right.get(tenure_col,''); pay=right.get(payout_col,'') if payout_col else ''; rt=right.get(rate_col,'')
            if ten and rt:
                line=f"Tenure: {ten}"
                if pay: line+=f", Payout: {pay}"
                line+=f", Profit Rate: {rt}"
                profit_lines.append(line)

    seen=set(); clean=[]
    for f in feature_lines:
        if f not in seen: seen.add(f); clean.append(f)

    parts=[]
    if clean:        parts.append("\n".join(f"- {f}" for f in clean))
    if profit_lines: parts.append("Profit Rate Information:\n"+"\n".join(f"- {p}" for p in profit_lines))
    return "\n\n".join(parts)

def convert_plan_table(block_rows, question, account_name):
    if not block_rows: return ""
    first=col_vals(block_rows[0]); headers=[v for _,v in first if not is_noise(v)]
    lines=[f"Available Plans ({' | '.join(headers[1:])}):" if len(headers)>1 else "Available Plans:"]
    for cd in block_rows[1:]:
        vals=[v for _,v in col_vals(cd) if not is_noise(v)]
        if not vals: continue
        name=re.sub(r'^[o•\-]\s*','',vals[0]).strip()
        rest=vals[1:]
        lines.append(f"  • {name}: {', '.join(rest)}" if rest else f"  • {name}")
    return "\n".join(lines)

def convert_generic(block_rows, question, account_name):
    lines=[]
    for cd in block_rows:
        vals=[v for v in just_vals(cd) if not is_noise(v)]
        if vals: lines.append("  • "+" | ".join(vals))
    return "\n".join(lines)

def detect_table_type(block_rows):
    if not block_rows: return None
    first=col_vals(block_rows[0]); vals=[v for _,v in first]
    if any(is_question(v) for v in vals) and len([(c,v) for c,v in first if not is_question(v) and not is_noise(v)])>=2:
        return 'multi_section'
    if any('profit payment' in v.lower() for v in vals) and any('profit rate' in v.lower() or 'rate' in v.lower() for v in vals):
        return 'profit_inline'
    if any('plan type' in v.lower() or 'premium' in v.lower() for v in vals):
        return 'plan_table'
    for cd in block_rows:
        if any(is_pct(v) for v in just_vals(cd)): return 'profit_inline'
    return 'generic'

def table_to_prose(block_rows, question, account_name):
    if not block_rows: return ""
    t=detect_table_type(block_rows)
    if t=='multi_section':  return convert_multi_section(block_rows, question, account_name)
    if t=='profit_inline':  return convert_profit_inline(block_rows, question, account_name)
    if t=='plan_table':     return convert_plan_table(block_rows, question, account_name)
    return convert_generic(block_rows, question, account_name)

# ── Main extractor ───────────────────────────────────────────────────────────

def extract_qa(ws, sheet_name, account_name):
    data=read_sparse(ws); all_rows=rows_list(data)
    qa_pairs=[]; i=0

    while i<len(all_rows):
        row_num,cd=all_rows[i]
        cvs=col_vals(cd); all_vals=[v for _,v in cvs]
        row_text=" ".join(all_vals); lower=row_text.lower().strip()

        if lower in {"main",sheet_name.lower(),account_name.lower()}: i+=1; continue

        question_val=next((v for v in all_vals if is_question(v)),None)
        if not question_val: i+=1; continue

        non_q_cells=[(c,v) for c,v in cvs if not is_question(v) and not is_noise(v)]
        has_inline_headers=len(non_q_cells)>=2

        if has_inline_headers:
            question=question_val
            table_block=[cd]; i+=1
            while i<len(all_rows):
                _,nc=all_rows[i]; nv=just_vals(nc); nt=" ".join(nv)
                if not nv: i+=1; continue
                if any(is_question(v) for v in nv): break
                table_block.append(nc); i+=1
            answer=table_to_prose(table_block,question,account_name)
            if question and answer: qa_pairs.append((question,answer))
        else:
            question=question_val; i+=1
            plain=[]; table_block=[]
            while i<len(all_rows):
                _,nc=all_rows[i]; nv=just_vals(nc); nt=" ".join(nv)
                if not nv: i+=1; continue
                if any(is_question(v) for v in nv): break
                if is_table_row(nc): table_block.append(nc)
                else: plain.append(nt)
                i+=1
            parts=[]
            if plain:       parts.append(" ".join(plain))
            if table_block: parts.append(table_to_prose(table_block,question,account_name))
            answer="\n\n".join(p for p in parts if p).strip()
            if question and answer: qa_pairs.append((question,answer))

    return qa_pairs

# ── Rate sheet ───────────────────────────────────────────────────────────────

def parse_rate_sheet(wb):
    ws_name=next((s for s in wb.sheetnames if "rate" in s.lower()),None)
    if not ws_name: return []
    ws=wb[ws_name]; data=read_sparse(ws); rows=rows_list(data)
    records=[]; seen=set(); current_account=""
    for _,cd in rows:
        vals=just_vals(cd); ne=[v for v in vals if not is_noise(v)]
        if not ne: continue
        if len(ne)==1 and not is_pct(ne[0]): current_account=ne[0]; continue
        if not any(is_pct(v) for v in ne): continue
        q=f"What is the profit rate for {current_account}?" if current_account else f"Profit rate: {ne[0]}"
        a=" | ".join(ne); content=anonymize(f"[Rate Sheet]\nQuestion: {q}\nAnswer: {a}")
        h=sha(content)
        if h not in seen:
            seen.add(h)
            records.append({"hash_id":h,"metadata":{"source":EXCEL_PATH,"sheet":ws_name,"account_name":current_account,"question":q},"content":content,"structured_data":{"question":q,"answer":a,"account":current_account}})
    return records

# ── Entry point ──────────────────────────────────────────────────────────────

def process_excel():
    print(f"Opening {EXCEL_PATH} ...")
    wb=load_workbook(EXCEL_PATH,data_only=True)
    all_records=[]; seen_hashes=set()

    for sheet_name in wb.sheetnames:
        if sheet_name in SKIP_SHEETS: print(f"  Skipping: {sheet_name}"); continue
        account_name=ACCOUNT_NAME_MAP.get(sheet_name,DEFAULT_NAME)
        print(f"  Processing: {sheet_name} → {account_name}")
        ws=wb[sheet_name]; qa_pairs=extract_qa(ws,sheet_name,account_name)
        if not qa_pairs: print(f"    ⚠  No Q&A pairs found"); continue
        print(f"    ✓  {len(qa_pairs)} pairs")
        for q,a in qa_pairs:
            # ── KEY CHANGE: repeat question keywords in content for better retrieval ──
            enriched=anonymize(
                f"[Account: {account_name}]\n"
                f"[Sheet: {sheet_name}]\n"
                f"Question: {q}\n"
                f"Answer: {a}\n"
                f"[Keywords: {account_name} {q}]"
            )
            h=sha(enriched)
            if h in seen_hashes: continue
            seen_hashes.add(h)
            all_records.append({
                "hash_id":h,
                "metadata":{"source":EXCEL_PATH,"sheet":sheet_name,"account_name":account_name,"question":q},
                "content":enriched,
                "structured_data":{"question":q,"answer":a,"account":account_name}
            })

    print("\n  Parsing Rate Sheet...")
    for rec in parse_rate_sheet(wb):
        h=rec["hash_id"]
        if h not in seen_hashes: seen_hashes.add(h); all_records.append(rec)

    if Path(RATE_SHEET_JSONL).exists():
        with open(RATE_SHEET_JSONL,encoding="utf-8") as f:
            for line in f:
                rec=json.loads(line); h=rec.get("hash_id","")
                if h and h not in seen_hashes: seen_hashes.add(h); all_records.append(rec)

    Path(OUTPUT_JSONL).parent.mkdir(parents=True,exist_ok=True)
    with open(OUTPUT_JSONL,"w",encoding="utf-8") as f:
        for rec in all_records: f.write(json.dumps(rec,ensure_ascii=False)+"\n")

    print(f"\n✅ Done. {len(all_records)} total records → {OUTPUT_JSONL}")

    print("\n─── NWA sample ───")
    for rec in all_records:
        if rec['metadata'].get('sheet')=='NWA':
            print(f"\nQ: {rec['structured_data']['question'][:80]}")
            print(f"A:\n{rec['structured_data']['answer'][:600]}")
            print("---")

if __name__=="__main__":
    process_excel()